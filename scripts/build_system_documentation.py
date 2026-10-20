from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SOURCE = DOCS / "SYSTEM_DOCUMENTATION.md"
OUTPUT = DOCS / "export" / "PortGen_System_Documentation.docx"
CROPS = ROOT / "tmp" / "documentation-screenshot-crops"
SCREENSHOT_ROOT = Path(os.environ.get("PORTGEN_DOC_SCREENSHOT_ROOT", DOCS / "screenshots"))

TRACKER_GROUPS = {
    "S": "Shneiderman's 8 Golden Rules",
    "N": "Norman's 7 Principles",
    "H": "Nielsen's 10 Usability Heuristics",
    "U": "Universal Design",
    "C": "Color theory",
    "V": "Visual design principles",
    "R": "User needs",
    "A": "Authentication usability",
}

SCREENSHOTS = [
    ("step-1-your-info", "Create form, step 1", "S8, N4, H5, U7, V3"),
    ("step-2-choose-template", "Template selection, step 2", "N2, H6, C1, V3"),
    ("step-3-done-share", "Completion and sharing, step 3", "S3, H1, C4, V7; the signed owner link is redacted in saved evidence"),
    ("template-modern", "Modern portfolio preview", "H6, C1, C3, V5"),
    ("template-simple", "Public Simple portfolio", "H8, C1, V2, V4"),
    ("template-creative", "Public Creative portfolio", "H6, C1, C2, V6"),
    ("manage", "Private portfolio management", "S7, H1, H3, H4"),
    ("edit", "Edit portfolio", "H3, N4, U2, V2"),
    ("delete-undo", "Delete and undo", "S6, H3, U5; ten-second undo state"),
    ("help-faq", "Help and FAQ", "N1, H10, U1"),
    ("keyboard-shortcuts", "Keyboard shortcuts dialog", "S2, H7, U2"),
    ("error-403", "Invalid private link, 403", "S3, H9, C4"),
    ("error-404", "Missing portfolio, 404", "S3, H9, C4"),
    ("error-419", "Expired form, 419", "S3, N6, H9"),
    ("error-429", "Rate limited, 429", "S3, N6, H9"),
    ("error-500", "Server error, 500", "S3, N6, H9"),
    ("dark-mode-creative", "Creative portfolio in dark mode", "C2, C3, V5"),
]


def parse_pipe_table(path: Path, expected: int) -> list[list[str]]:
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if len(cells) == expected and not all(set(cell) <= {"-", ":"} for cell in cells):
            rows.append(cells)
    return rows


def markdown_table(headers: list[str], rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in rows)
    return "\n".join(lines)


def tracker_block(first_table: int) -> str:
    rows = parse_pipe_table(DOCS / "HCI_TRACKER.md", 6)
    by_group = {prefix: [] for prefix in TRACKER_GROUPS}
    for row in rows:
        if row[0] in {"ID"}:
            continue
        by_group[row[0][0]].append(row)
    blocks = ["<!-- LANDSCAPE_START -->"]
    number = first_table
    for prefix, title in TRACKER_GROUPS.items():
        blocks.extend([
            f"#### {title}",
            f"**Table {number}. {title} HCI tracker**",
            markdown_table(["ID", "Principle", "Where applied", "How applied", "Evidence", "Status"], by_group[prefix]),
            "",
        ])
        number += 1
    blocks.append("<!-- LANDSCAPE_END -->")
    return "\n".join(blocks)


def screenshot_gallery(first_figure: int) -> str:
    blocks = ["<!-- PAGEBREAK -->"]
    figure = first_figure
    names = {row[0]: row[1] for row in parse_pipe_table(DOCS / "HCI_TRACKER.md", 6) if row[0] not in {"ID"}}
    for name, caption, principles in SCREENSHOTS:
        ids = [value.strip() for value in principles.split(";", 1)[0].split(",")]
        annotation = "; ".join(f"{code} {names.get(code, 'Principle')}" for code in ids)
        mobile = f"![Mobile screenshot, {caption}](screenshots/mobile/{name}.png)"
        desktop = f"![Desktop screenshot, {caption}](screenshots/desktop/{name}.png)"
        blocks.extend([
            f"**Figure {figure}. {caption}**",
            "| Mobile view | Desktop view |",
            "|---|---|",
            f"| {mobile} | {desktop} |",
            f"**HCI principles:** {annotation}. {principles.split(';', 1)[1].strip() if ';' in principles else ''}".rstrip(),
            "<!-- PAGEBREAK -->",
        ])
        figure += 1
    return "\n".join(blocks)


def decisions_block() -> str:
    rows = parse_pipe_table(DOCS / "DECISIONS.md", 3)
    rows = [row for row in rows if row[0] != "Date"]
    return markdown_table(["Date", "Decision", "Reason"], rows)


def expand_source() -> str:
    source = SOURCE.read_text(encoding="utf-8")
    if "<!-- HCI_TRACKER_TABLES -->" in source:
        source = source.replace("<!-- HCI_TRACKER_TABLES -->", tracker_block(6))
    if "<!-- HCI_TRACKER_APPENDIX -->" in source:
        source = source.replace("<!-- HCI_TRACKER_APPENDIX -->", tracker_block(18))
    if "<!-- SCREENSHOT_GALLERY -->" in source:
        source = source.replace("<!-- SCREENSHOT_GALLERY -->", screenshot_gallery(21))
    if "<!-- DECISIONS_TABLE -->" in source:
        source = source.replace("<!-- DECISIONS_TABLE -->", decisions_block())
    SOURCE.write_text(source, encoding="utf-8")
    return source


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def set_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("PortGen System Documentation - Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0, 0, 0)
    add_page_field(paragraph)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "174F78")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*|\[[^\]]+\]\([^)]+\))")
IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        elif token.startswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def set_run_font(run, size=9):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


def image_path(relative: str) -> Path:
    path = Path(relative)
    if path.parts and path.parts[0] == "screenshots":
        return SCREENSHOT_ROOT.joinpath(*path.parts[1:]).resolve()
    return (DOCS / path).resolve()


def screenshot_viewport_path(relative: str, mode: str) -> Path:
    source = image_path(relative)
    target = CROPS / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    viewport = (375, 812) if mode == "mobile" else (1280, 800)
    with Image.open(source) as capture:
        width = min(capture.width, viewport[0])
        height = min(capture.height, viewport[1])
        capture.crop((0, 0, width, height)).save(target)
    return target


def table_widths(headers: list[str], image_row=False) -> list[float]:
    if image_row:
        return [2.0, 5.05]
    if len(headers) == 6 and headers[0] == "ID":
        return [0.42, 1.25, 1.8, 2.25, 3.45, 0.55]
    if len(headers) == 4 and headers[0] == "Column":
        return [1.0, 1.2, 1.7, 3.15]
    if len(headers) == 4 and headers[0] == "Method":
        return [0.8, 2.5, 2.55, 1.2]
    if len(headers) == 3 and headers[0] == "Date":
        return [1.0, 2.55, 4.3]
    if len(headers) == 3 and headers[0] == "Persona":
        return [1.0, 2.2, 4.65]
    return [7.05 / max(len(headers), 1)] * len(headers)


def mark_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True


def mark_row_no_split(row) -> None:
    no_split = OxmlElement("w:cantSplit")
    row._tr.get_or_add_trPr().append(no_split)


def format_cell(cell, header=False, body_size=7.4):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    if header:
        shade(cell, "123D65")
    for paragraph in cell.paragraphs:
        compact_tracker = not header and body_size < 7
        paragraph.paragraph_format.space_after = Pt(1 if compact_tracker else 2)
        paragraph.paragraph_format.space_before = Pt(0 if compact_tracker else 1)
        if compact_tracker:
            paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, 7.2 if header else body_size)
            if header:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)


def add_markdown_table(doc, raw_rows: list[list[str]]):
    headers = raw_rows[0]
    body_rows = raw_rows[1:]
    widths = table_widths(headers)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    mark_header(table.rows[0])
    for i, heading in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.text = heading
        format_cell(cell, header=True)
    for row_index, values in enumerate(body_rows):
        row = table.add_row()
        mark_row_no_split(row)
        cells = row.cells
        for i, value in enumerate(values):
            cells[i].width = Inches(widths[i])
            if row_index % 2 == 1:
                shade(cells[i], "EFF4F8")
            match = IMAGE.fullmatch(value.strip())
            if match:
                alt, relative = match.groups()
                paragraph = cells[i].paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                mode = "mobile" if "/mobile/" in relative else "desktop"
                label = paragraph.add_run("Mobile - 375 x 812" if mode == "mobile" else "Desktop - 1280 x 800")
                set_run_font(label, 7.2)
                paragraph.add_run().add_break()
                picture = paragraph.add_run().add_picture(str(screenshot_viewport_path(relative, mode)), width=Inches(2.0 if mode == "mobile" else 3.2))
                picture._inline.docPr.set("descr", alt)
                picture._inline.docPr.set("title", alt)
            else:
                cells[i].text = ""
                paragraph = cells[i].paragraphs[0]
                add_inline(paragraph, value)
            format_cell(cells[i], body_size=6.3 if headers and headers[0] == "ID" else 7.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def set_section(section, landscape: bool):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width, section.page_height = Inches(11), Inches(8.5)
        section.left_margin = section.right_margin = Inches(0.55)
        section.top_margin = section.bottom_margin = Inches(0.6)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(0.68)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.72)
    set_footer(section)


def build_docx(source: str) -> None:
    doc = Document()
    for section in doc.sections:
        set_section(section, landscape=False)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    title_ppr = styles["Title"]._element.pPr
    title_border = title_ppr.find(qn("w:pBdr")) if title_ppr is not None else None
    if title_border is not None:
        title_ppr.remove(title_border)

    lines = source.splitlines()
    index = 0
    landscape = False
    pending_list = None
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == "<!-- LANDSCAPE_START -->":
            section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            set_section(section, landscape=True)
            landscape = True
            index += 1
            continue
        if line == "<!-- LANDSCAPE_END -->":
            if any(candidate.strip() for candidate in lines[index + 1:]):
                section = doc.add_section(WD_SECTION_START.NEW_PAGE)
                set_section(section, landscape=False)
            landscape = False
            index += 1
            continue
        if line == "<!-- PAGEBREAK -->":
            doc.add_page_break()
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|?\s*:?-{3,}", lines[index + 1].strip()):
            raw_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(set(cell) <= {"-", ":"} for cell in row):
                    raw_rows.append(row)
                index += 1
            if raw_rows:
                add_markdown_table(doc, raw_rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if text in {"2. Users and needs", "12. Appendices"}:
                doc.add_page_break()
            if text in TRACKER_GROUPS.values() and text != TRACKER_GROUPS["S"]:
                doc.add_page_break()
            style = "Title" if level == 1 else f"Heading {min(level - 1, 3)}"
            paragraph = doc.add_paragraph(style=style)
            add_inline(paragraph, text)
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue
        if line.startswith("---"):
            index += 1
            continue
        if line.startswith("!["):
            match = IMAGE.fullmatch(line)
            if match:
                alt, relative = match.groups()
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                normalized_relative = relative.replace("\\", "/")
                simple_review_image = "screenshots/simple/after/" in normalized_relative
                mode = "mobile" if simple_review_image and "375-" in relative else "desktop"
                width = 3.25 if "system-architecture" in relative else 5.0 if "portfolio-schema" in relative else 2.0 if simple_review_image and mode == "mobile" else 6.2
                source = screenshot_viewport_path(relative, mode) if simple_review_image else image_path(relative)
                picture = paragraph.add_run().add_picture(str(source), width=Inches(width))
                picture._inline.docPr.set("descr", alt)
                picture._inline.docPr.set("title", alt)
            index += 1
            continue
        item = re.match(r"^(\s*)(?:- |\d+\. )(.*)$", line)
        if item:
            number = re.match(r"^\s*(\d+)\. ", line)
            if number:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.27)
                paragraph.paragraph_format.first_line_indent = Inches(-0.25)
                paragraph.add_run(f"{number.group(1)}. ")
                add_inline(paragraph, re.sub(r"^\s*\d+\. ", "", line))
            else:
                paragraph = doc.add_paragraph(style="List Bullet")
                add_inline(paragraph, item.group(2))
            index += 1
            continue
        if line.startswith("<!--") and line.endswith("-->"):
            index += 1
            continue
        paragraph = doc.add_paragraph()
        add_inline(paragraph, line.replace("  ", " "))
        if line.startswith("**Table ") or line.startswith("**Figure "):
            paragraph.paragraph_format.keep_with_next = True
        index += 1

    for section in doc.sections:
        set_footer(section)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


def main() -> None:
    expanded = expand_source()
    build_docx(expanded)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
